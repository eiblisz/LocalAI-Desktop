import json
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .config import OUTPUT_DIR

RED = "B92F3B"
DARK_RED = "8E1F2D"
LIGHT_GREY = "F4F6F8"
WHITE = "FFFFFF"
TEXT = "1F2630"


def _clean_json(text):
    stripped = text.strip()
    fence = chr(96) * 3
    if stripped.startswith(fence):
        stripped = stripped[len(fence):].lstrip()
        if stripped.lower().startswith("json"):
            stripped = stripped[4:].lstrip()
    if stripped.endswith(fence):
        stripped = stripped[:-len(fence)].rstrip()
    return json.loads(stripped)


def _safe_sheet_name(name):
    clean = re.sub(r"[:\\/?*\[\]]", "-", str(name)).strip() or "Sheet"
    return clean[:31]


def _style_sheet(ws, title, columns=1):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"
    end_col = max(1, columns)
    ws.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=end_col,
    )
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial", size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=DARK_RED)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30


def _style_table(ws, header_row=3):
    thin = Side(style="thin", color="D9DEE5")
    for cell in ws[header_row]:
        if cell.value is not None:
            cell.font = Font(name="Arial", size=11, bold=True, color=WHITE)
            cell.fill = PatternFill("solid", fgColor=RED)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            cell.border = Border(bottom=thin)

    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.font = Font(name="Arial", size=10.5, color=TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(
                bottom=Side(style="hair", color="E6E9ED")
            )


def _auto_width(ws, max_width=48):
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        width = 10
        for cell in ws[letter]:
            if cell.value is not None:
                width = max(
                    width,
                    min(len(str(cell.value)) + 2, max_width),
                )
        ws.column_dimensions[letter].width = width


def create_conversation_excel(
    messages,
    title="Local AI Conversation",
    output_dir=OUTPUT_DIR,
):
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / (
        "local_ai_workbook_"
        + datetime.now().strftime("%Y%m%d_%H%M%S")
        + ".xlsx"
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Conversation"
    _style_sheet(ws, title, columns=3)
    ws.append([])
    ws.append(["Role", "Content", "Artifact path"])

    for message in messages:
        role = message.get("role", "")
        if role == "system":
            continue
        ws.append([
            role.upper(),
            message.get("content", ""),
            message.get("path", "") if role == "artifact" else "",
        ])

    _style_table(ws)
    _auto_width(ws)
    if ws.max_row >= 3:
        ws.auto_filter.ref = f"A3:C{ws.max_row}"
    wb.save(path)
    return path


def create_structured_excel(
    model_output,
    title="Local AI Workbook",
    output_dir=OUTPUT_DIR,
):
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / (
        "local_ai_workbook_"
        + datetime.now().strftime("%Y%m%d_%H%M%S")
        + ".xlsx"
    )

    try:
        payload = _clean_json(model_output)
    except Exception:
        payload = {
            "title": title,
            "sheets": [{
                "name": "Content",
                "headers": ["Content"],
                "rows": [[model_output]],
            }],
        }

    workbook_title = str(payload.get("title") or title)
    sheets = payload.get("sheets") or []
    if not sheets:
        sheets = [{
            "name": "Content",
            "headers": ["Content"],
            "rows": [["No data"]],
        }]

    wb = Workbook()
    wb.remove(wb.active)

    for index, spec in enumerate(sheets):
        name = _safe_sheet_name(
            spec.get("name") or f"Sheet {index + 1}"
        )
        ws = wb.create_sheet(name)
        headers = list(spec.get("headers") or ["Value"])
        rows = spec.get("rows") or []
        _style_sheet(ws, workbook_title, columns=len(headers))
        ws.append([])
        ws.append(headers)

        for row in rows:
            if isinstance(row, list):
                ws.append(row)
            else:
                ws.append([row])

        _style_table(ws)
        _auto_width(ws)
        if ws.max_column and ws.max_row >= 3:
            end = get_column_letter(ws.max_column)
            ws.auto_filter.ref = f"A3:{end}{ws.max_row}"

    wb.save(path)
    return path
