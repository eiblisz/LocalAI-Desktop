from __future__ import annotations

import csv
import html
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from urllib.parse import urlparse

import markdown
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from PySide6.QtCore import Qt, QUrl, Signal
from PySide6.QtGui import QDesktopServices, QPixmap
from .browser_navigation_authority import (
    ACTION_EXTERNAL,
    ACTION_INTERNAL_SAME,
    ACTION_INTERNAL_TAB,
    BrowserNavigationAuthority,
)

from PySide6.QtWidgets import (
    QFileDialog,
    QFrame,
    QHBoxLayout,
    QLabel,
    QLineEdit,
    QMessageBox,
    QPushButton,
    QScrollArea,
    QTabWidget,
    QTableWidget,
    QTableWidgetItem,
    QTextBrowser,
    QVBoxLayout,
    QWidget,
)

try:
    from PySide6.QtWebEngineCore import QWebEnginePage
    from PySide6.QtWebEngineWidgets import QWebEngineView
except Exception:  # pragma: no cover - optional runtime fallback
    QWebEnginePage = None
    QWebEngineView = None

try:
    from PySide6.QtPdf import QPdfDocument
    from PySide6.QtPdfWidgets import QPdfView
except Exception:  # pragma: no cover - optional runtime fallback
    QPdfDocument = None
    QPdfView = None


if QWebEnginePage is not None:
    class _AuthorityWebPage(QWebEnginePage):
        def __init__(self, authority, open_resource=None, parent=None):
            super().__init__(parent)
            self.authority = authority
            self.open_resource = open_resource

        def acceptNavigationRequest(self, url, navigation_type, is_main_frame):
            if not is_main_frame:
                return True

            decision = self.authority.decide(
                url.toString(),
                source="page_link",
            )
            if decision.action == ACTION_INTERNAL_SAME:
                return True
            if (
                decision.action == ACTION_INTERNAL_TAB
                and callable(self.open_resource)
            ):
                self.open_resource(decision.target)
            return False
else:  # pragma: no cover - optional runtime fallback
    _AuthorityWebPage = None


RESOURCE_IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp"}
RESOURCE_TEXT_SUFFIXES = {
    ".txt",
    ".log",
    ".py",
    ".json",
    ".toml",
    ".yaml",
    ".yml",
    ".ini",
    ".cfg",
}
RESOURCE_SHEET_SUFFIXES = {".xlsx", ".xlsm", ".csv"}


def classify_resource(target) -> str:
    value = str(target or "").strip()
    parsed = urlparse(value)
    if parsed.scheme.lower() in {"http", "https"}:
        return "browser"

    path = Path(target)
    suffix = path.suffix.lower()
    if suffix in {".html", ".htm"}:
        return "html"
    if suffix == ".pdf":
        return "pdf"
    if suffix == ".docx":
        return "docx"
    if suffix in RESOURCE_SHEET_SUFFIXES:
        return "spreadsheet"
    if suffix == ".md":
        return "markdown"
    if suffix in RESOURCE_IMAGE_SUFFIXES:
        return "image"
    if suffix in RESOURCE_TEXT_SUFFIXES:
        return "text"
    return "unknown"


def resource_title(target) -> str:
    value = str(target or "").strip()
    parsed = urlparse(value)
    if parsed.scheme.lower() in {"http", "https"}:
        host = parsed.netloc or "Browser"
        tail = Path(parsed.path).name
        return (tail or host)[:64]

    path = Path(target)
    return (path.name or "Resource")[:64]


def resource_identity(target) -> str:
    value = str(target or "").strip()
    parsed = urlparse(value)
    if parsed.scheme.lower() in {"http", "https"}:
        return value
    return str(Path(value).resolve())


def save_resource_copy(source, destination):
    source_path = Path(source).resolve()
    destination_path = Path(destination).expanduser().resolve()
    if not source_path.exists() or not source_path.is_file():
        raise FileNotFoundError(str(source_path))
    destination_path.parent.mkdir(parents=True, exist_ok=True)
    if source_path == destination_path:
        return destination_path
    shutil.copy2(source_path, destination_path)
    return destination_path


def _default_save_path(path):
    path = Path(path).resolve()
    downloads = Path.home() / "Downloads"
    base = downloads if downloads.exists() else Path.home()
    return base / path.name


def _save_resource_as_dialog(parent, path):
    source = Path(path).resolve()
    destination, _selected_filter = QFileDialog.getSaveFileName(
        parent,
        "Save As",
        str(_default_save_path(source)),
        f"{source.suffix.upper().lstrip('.')} files (*{source.suffix});;All files (*.*)",
    )
    if not destination:
        return None
    try:
        return save_resource_copy(source, destination)
    except Exception as exc:
        QMessageBox.critical(parent, "Save file error", str(exc))
        return None


def _add_save_as_toolbar(layout, path, parent):
    toolbar = QFrame(parent)
    controls = QHBoxLayout(toolbar)
    controls.setContentsMargins(8, 7, 8, 7)
    controls.addStretch()
    save_button = QPushButton("SAVE AS")
    save_button.setToolTip("Save a copy of this file to another location.")
    save_button.clicked.connect(
        lambda _checked=False, source=Path(path).resolve(): _save_resource_as_dialog(
            parent,
            source,
        )
    )
    controls.addWidget(save_button)
    layout.addWidget(toolbar)
    return save_button


def find_libreoffice_executable():
    for name in ("soffice", "libreoffice"):
        executable = shutil.which(name)
        if executable:
            return Path(executable)

    candidates = []
    for env_name in ("ProgramFiles", "ProgramFiles(x86)"):
        root = os.environ.get(env_name)
        if root:
            candidates.append(Path(root) / "LibreOffice" / "program" / "soffice.exe")

    for candidate in candidates:
        if candidate.exists() and candidate.is_file():
            return candidate
    return None


def convert_docx_to_pdf_preview(path, output_dir, executable=None, timeout=45):
    source = Path(path).resolve()
    destination_dir = Path(output_dir).resolve()
    if not source.exists() or not source.is_file():
        raise FileNotFoundError(str(source))

    executable = Path(executable) if executable else find_libreoffice_executable()
    if executable is None:
        return None

    destination_dir.mkdir(parents=True, exist_ok=True)
    kwargs = {
        "args": [
            str(executable),
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(destination_dir),
            str(source),
        ],
        "capture_output": True,
        "text": True,
        "timeout": timeout,
        "check": False,
    }
    if hasattr(subprocess, "CREATE_NO_WINDOW"):
        kwargs["creationflags"] = subprocess.CREATE_NO_WINDOW

    completed = subprocess.run(**kwargs)
    preview = destination_dir / f"{source.stem}.pdf"
    if completed.returncode != 0 or not preview.exists():
        return None
    return preview


def render_docx_html(path) -> str:
    document = Document(str(path))
    parts = [
        "<div style='font-family: Segoe UI; font-size:15px; color:#20242A;"
        "line-height:1.55; padding:28px;'>"
    ]

    for paragraph in document.paragraphs:
        text = html.escape(paragraph.text)
        if not text.strip():
            parts.append("<div style='height:10px;'></div>")
            continue

        style_name = str(getattr(paragraph.style, "name", "") or "")
        match = re.match(r"Heading\s+(\d+)", style_name, flags=re.IGNORECASE)
        if match:
            level = min(max(int(match.group(1)), 1), 6)
            parts.append(f"<h{level}>{text}</h{level}>")
        else:
            parts.append(f"<p>{text}</p>")

    for table in document.tables:
        parts.append(
            "<table style='border-collapse:collapse;width:100%;margin:16px 0;'>"
        )
        for row in table.rows:
            parts.append("<tr>")
            for cell in row.cells:
                value = html.escape(cell.text)
                parts.append(
                    "<td style='border:1px solid #B9C0C8;padding:7px;"
                    f"vertical-align:top;'>{value}</td>"
                )
            parts.append("</tr>")
        parts.append("</table>")

    parts.append("</div>")
    return "".join(parts)


def spreadsheet_preview(path, max_rows=500, max_columns=50):
    path = Path(path)
    suffix = path.suffix.lower()

    if suffix == ".csv":
        rows = []
        truncated = False
        with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
            reader = csv.reader(handle)
            for index, row in enumerate(reader):
                if index >= max_rows:
                    truncated = True
                    break
                rows.append([str(value) for value in row[:max_columns]])
                if len(row) > max_columns:
                    truncated = True
        return [
            {
                "name": path.stem or "CSV",
                "rows": rows,
                "truncated": truncated,
            }
        ]

    workbook = load_workbook(
        filename=str(path),
        read_only=True,
        data_only=False,
    )
    previews = []
    try:
        for worksheet in workbook.worksheets:
            rows = []
            truncated = False
            for row_index, row in enumerate(
                worksheet.iter_rows(values_only=True),
                start=1,
            ):
                if row_index > max_rows:
                    truncated = True
                    break
                values = list(row)
                if len(values) > max_columns:
                    values = values[:max_columns]
                    truncated = True
                rows.append([
                    "" if value is None else str(value)
                    for value in values
                ])
            previews.append(
                {
                    "name": worksheet.title,
                    "rows": rows,
                    "truncated": truncated,
                }
            )
    finally:
        workbook.close()

    return previews


def pdf_text_fallback(path, max_pages=30):
    reader = PdfReader(str(path))
    parts = []
    for index, page in enumerate(reader.pages[:max_pages], start=1):
        text = page.extract_text() or ""
        parts.append(f"--- PAGE {index} ---\n{text}")
    if len(reader.pages) > max_pages:
        parts.append("[PDF preview truncated]")
    return "\n\n".join(parts)


class BrowserView(QWidget):
    title_changed = Signal(str)

    def __init__(
        self,
        url,
        open_resource=None,
        navigation_authority=None,
        parent=None,
    ):
        super().__init__(parent)
        self.open_resource = open_resource
        self.navigation_authority = (
            navigation_authority or BrowserNavigationAuthority()
        )
        self._initial_url = str(url or "").strip()

        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(0)

        toolbar = QFrame()
        controls = QHBoxLayout(toolbar)
        controls.setContentsMargins(8, 7, 8, 7)

        self.back_button = QPushButton("←")
        self.forward_button = QPushButton("→")
        self.reload_button = QPushButton("↻")
        self.address = QLineEdit()
        self.external_button = QPushButton("External")

        controls.addWidget(self.back_button)
        controls.addWidget(self.forward_button)
        controls.addWidget(self.reload_button)
        controls.addWidget(self.address, 1)
        controls.addWidget(self.external_button)
        root.addWidget(toolbar)

        if QWebEngineView is not None:
            self.view = QWebEngineView()
            if _AuthorityWebPage is not None:
                self.view.setPage(
                    _AuthorityWebPage(
                        self.navigation_authority,
                        open_resource=self.open_resource,
                        parent=self.view,
                    )
                )
            self.back_button.clicked.connect(self.view.back)
            self.forward_button.clicked.connect(self.view.forward)
            self.reload_button.clicked.connect(self.view.reload)
            self.address.returnPressed.connect(self._navigate_address)
            self.external_button.clicked.connect(self._open_external)
            self.view.urlChanged.connect(self._url_changed)
            self.view.titleChanged.connect(self._title_changed)

            page = self.view.page()
            if hasattr(page, "newWindowRequested"):
                page.newWindowRequested.connect(self._new_window_requested)

            root.addWidget(self.view, 1)
            self.load(self._initial_url)
        else:
            self.view = QTextBrowser()
            self.view.setOpenExternalLinks(False)
            self.view.setOpenLinks(False)
            self.view.anchorClicked.connect(self._fallback_link_clicked)
            self.back_button.clicked.connect(self.view.backward)
            self.forward_button.clicked.connect(self.view.forward)
            self.reload_button.clicked.connect(lambda: self.load(self.address.text()))
            self.address.returnPressed.connect(self._navigate_address)
            self.external_button.clicked.connect(self._open_external)
            root.addWidget(self.view, 1)
            self.load(self._initial_url)

    def load(self, value, source="initial_load"):
        decision = self.navigation_authority.decide(
            value,
            source=source,
        )
        if not decision.allowed:
            return False

        if decision.action == ACTION_EXTERNAL:
            return bool(QDesktopServices.openUrl(QUrl(decision.target)))

        if (
            decision.action == ACTION_INTERNAL_TAB
            and callable(self.open_resource)
        ):
            self.open_resource(decision.target)
            return True

        self.address.setText(decision.target)
        qurl = QUrl(decision.target)
        if not qurl.scheme():
            qurl = QUrl.fromLocalFile(str(Path(decision.target).resolve()))

        if QWebEngineView is not None and isinstance(self.view, QWebEngineView):
            self.view.setUrl(qurl)
        else:
            self.view.setSource(qurl)
        return True

    def _navigate_address(self):
        self.load(
            self.address.text().strip(),
            source="address_bar",
        )

    def _url_changed(self, qurl):
        self.address.setText(qurl.toString())

    def _title_changed(self, title):
        title = str(title or "").strip()
        if title:
            self.title_changed.emit(title[:64])

    def _new_window_requested(self, request):
        qurl = request.requestedUrl()
        if not qurl.isValid():
            return
        decision = self.navigation_authority.decide(
            qurl.toString(),
            source="new_window",
        )
        if (
            decision.action == ACTION_INTERNAL_TAB
            and callable(self.open_resource)
        ):
            self.open_resource(decision.target)
        elif decision.action == ACTION_INTERNAL_SAME:
            self.load(decision.target, source="page_link")

    def _fallback_link_clicked(self, qurl):
        decision = self.navigation_authority.decide(
            qurl.toString(),
            source="page_link",
        )
        if decision.action == ACTION_INTERNAL_SAME:
            self.load(decision.target, source="page_link")
        elif (
            decision.action == ACTION_INTERNAL_TAB
            and callable(self.open_resource)
        ):
            self.open_resource(decision.target)

    def _open_external(self):
        decision = self.navigation_authority.decide(
            self.address.text().strip(),
            source="external_button",
        )
        if decision.action == ACTION_EXTERNAL:
            QDesktopServices.openUrl(QUrl(decision.target))


class PdfViewWidget(QWidget):
    def __init__(
        self,
        path,
        parent=None,
        show_save_as=True,
        fit_full_page=False,
    ):
        super().__init__(parent)
        path = Path(path).resolve()

        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        if show_save_as:
            _add_save_as_toolbar(root, path, self)

        if QPdfDocument is not None and QPdfView is not None:
            self.document = QPdfDocument(self)
            error = self.document.load(str(path))
            self.view = QPdfView(self)
            self.view.setDocument(self.document)
            self._add_pdf_zoom_toolbar(root)
            try:
                self.view.setPageMode(QPdfView.PageMode.MultiPage)
                zoom_mode = (
                    QPdfView.ZoomMode.FitInView
                    if fit_full_page
                    else QPdfView.ZoomMode.FitToWidth
                )
                self.view.setZoomMode(zoom_mode)
                self._sync_zoom_label()
            except Exception:
                pass
            root.addWidget(self.view, 1)

            # QPdfDocument.Error.None_ is not available in every PySide6 build,
            # so only fall back when pageCount remains unusable after load.
            if self.document.pageCount() > 0:
                return

        fallback = QTextBrowser()
        fallback.setPlainText(pdf_text_fallback(path))
        root.addWidget(fallback, 1)


    def _add_pdf_zoom_toolbar(self, layout):
        toolbar = QFrame(self)
        controls = QHBoxLayout(toolbar)
        controls.setContentsMargins(8, 5, 8, 5)
        controls.addStretch()

        zoom_out = QPushButton("-")
        zoom_out.setToolTip("Zoom out")
        zoom_out.clicked.connect(lambda: self._zoom_by(0.9))

        self.zoom_label = QLabel("100%")
        self.zoom_label.setMinimumWidth(52)
        self.zoom_label.setAlignment(Qt.AlignmentFlag.AlignCenter)

        zoom_in = QPushButton("+")
        zoom_in.setToolTip("Zoom in")
        zoom_in.clicked.connect(lambda: self._zoom_by(1.1))

        reset = QPushButton("100%")
        reset.setToolTip("Reset zoom to 100%")
        reset.clicked.connect(lambda: self._set_zoom_percent(100))

        fit_width = QPushButton("FIT WIDTH")
        fit_width.clicked.connect(self._fit_width)

        fit_page = QPushButton("FIT PAGE")
        fit_page.clicked.connect(self._fit_page)

        for button in (zoom_out, zoom_in, reset, fit_width, fit_page):
            controls.addWidget(button)
        controls.insertWidget(controls.count() - 4, self.zoom_label)
        layout.addWidget(toolbar)

    def _set_zoom_percent(self, percent):
        percent = max(25, min(int(percent), 400))
        try:
            self.view.setZoomMode(QPdfView.ZoomMode.Custom)
        except Exception:
            pass
        self.view.setZoomFactor(percent / 100.0)
        self._sync_zoom_label()

    def _zoom_by(self, factor):
        current = max(0.25, min(float(self.view.zoomFactor()), 4.0))
        self._set_zoom_percent(round(current * factor * 100))

    def _fit_width(self):
        self.view.setZoomMode(QPdfView.ZoomMode.FitToWidth)
        self._sync_zoom_label()

    def _fit_page(self):
        self.view.setZoomMode(QPdfView.ZoomMode.FitInView)
        self._sync_zoom_label()

    def _sync_zoom_label(self):
        if not hasattr(self, "zoom_label"):
            return
        try:
            percent = round(float(self.view.zoomFactor()) * 100)
            self.zoom_label.setText(f"{percent}%")
        except Exception:
            self.zoom_label.setText("AUTO")


class DocumentView(QWidget):
    def __init__(self, path, parent=None):
        super().__init__(parent)
        path = Path(path).resolve()

        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        _add_save_as_toolbar(root, path, self)

        self._preview_temp = tempfile.TemporaryDirectory(
            prefix="localai-docx-preview-"
        )
        preview_pdf = None
        try:
            preview_pdf = convert_docx_to_pdf_preview(
                path,
                self._preview_temp.name,
            )
        except Exception:
            preview_pdf = None

        if preview_pdf is not None:
            status = QLabel("High-fidelity DOCX preview")
            status.setObjectName("muted")
            status.setStyleSheet("padding:4px 10px;color:#8F99A6;")
            root.addWidget(status)
            self.preview = PdfViewWidget(
                preview_pdf,
                parent=self,
                show_save_as=False,
                fit_full_page=True,
            )
            root.addWidget(self.preview, 1)
            return

        self._preview_temp.cleanup()
        self._preview_temp = None

        status = QLabel(
            "Basic DOCX preview - LibreOffice rendering is unavailable."
        )
        status.setObjectName("muted")
        status.setStyleSheet("padding:4px 10px;color:#8F99A6;")
        root.addWidget(status)

        view = QTextBrowser()
        view.setOpenExternalLinks(False)
        view.setStyleSheet(
            "QTextBrowser {"
            "background:#FFFFFF;"
            "color:#20242A;"
            "border:none;"
            "padding:0px;"
            "}"
        )
        view.setHtml(render_docx_html(path))
        root.addWidget(view, 1)


class SpreadsheetView(QWidget):
    def __init__(self, path, parent=None):
        super().__init__(parent)
        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        _add_save_as_toolbar(root, path, self)

        sheets = QTabWidget()
        previews = spreadsheet_preview(path)
        for preview in previews:
            rows = preview["rows"]
            column_count = max((len(row) for row in rows), default=0)
            table = QTableWidget(len(rows), column_count)
            table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
            table.setAlternatingRowColors(True)
            for row_index, row in enumerate(rows):
                for column_index, value in enumerate(row):
                    table.setItem(
                        row_index,
                        column_index,
                        QTableWidgetItem(value),
                    )
            table.resizeColumnsToContents()

            page = QWidget()
            page_layout = QVBoxLayout(page)
            page_layout.setContentsMargins(0, 0, 0, 0)
            page_layout.addWidget(table, 1)
            if preview.get("truncated"):
                note = QLabel(
                    "Preview limited to the first 500 rows and 50 columns."
                )
                note.setStyleSheet("color:#9099A6;padding:6px;")
                page_layout.addWidget(note)

            sheets.addTab(page, str(preview["name"])[:40])

        root.addWidget(sheets, 1)


class MarkdownTextView(QWidget):
    def __init__(self, path, markdown_mode=False, parent=None):
        super().__init__(parent)
        path = Path(path)
        text = path.read_text(encoding="utf-8", errors="replace")

        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)
        _add_save_as_toolbar(root, path, self)

        view = QTextBrowser()
        if markdown_mode:
            view.setHtml(
                markdown.markdown(
                    html.escape(text),
                    extensions=["fenced_code", "tables", "sane_lists", "nl2br"],
                )
            )
        else:
            view.setPlainText(text)
        root.addWidget(view, 1)


class ImageView(QWidget):
    def __init__(self, path, parent=None):
        super().__init__(parent)
        root = QVBoxLayout(self)
        root.setContentsMargins(0, 0, 0, 0)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        label = QLabel()
        label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        pixmap = QPixmap(str(Path(path).resolve()))
        label.setPixmap(pixmap)
        label.setScaledContents(False)
        scroll.setWidget(label)
        root.addWidget(scroll, 1)


class UnsupportedView(QWidget):
    def __init__(self, target, parent=None):
        super().__init__(parent)
        root = QVBoxLayout(self)
        message = QLabel(
            "This resource type does not have an internal preview yet.\n\n"
            f"{target}"
        )
        message.setWordWrap(True)
        message.setAlignment(Qt.AlignmentFlag.AlignCenter)
        root.addWidget(message, 1)


def create_resource_view(
    target,
    open_resource=None,
    navigation_authority=None,
    parent=None,
):
    kind = classify_resource(target)

    if kind == "browser":
        return BrowserView(
            target,
            open_resource=open_resource,
            navigation_authority=navigation_authority,
            parent=parent,
        )

    path = Path(target).resolve()
    if not path.exists() or not path.is_file():
        raise FileNotFoundError(str(path))

    if kind == "html":
        return BrowserView(
            path.as_uri(),
            open_resource=open_resource,
            navigation_authority=navigation_authority,
            parent=parent,
        )
    if kind == "pdf":
        return PdfViewWidget(path, parent=parent)
    if kind == "docx":
        return DocumentView(path, parent=parent)
    if kind == "spreadsheet":
        return SpreadsheetView(path, parent=parent)
    if kind == "markdown":
        return MarkdownTextView(path, markdown_mode=True, parent=parent)
    if kind == "text":
        return MarkdownTextView(path, markdown_mode=False, parent=parent)
    if kind == "image":
        return ImageView(path, parent=parent)
    return UnsupportedView(path, parent=parent)
