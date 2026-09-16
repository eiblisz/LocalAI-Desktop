import re
from datetime import datetime

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .config import OUTPUT_DIR

RED = "B92F3B"
DARK_RED = "8E1F2D"
LIGHT_RED = "F9EDEF"
TEXT = "1F2630"
MUTED = "6C737F"
LIGHT_GREY = "F4F6F8"


def _visible_messages(messages):
    return [
        m for m in messages
        if m.get("role", "").lower() not in {"system", "artifact"}
    ]


def _document_text(messages):
    visible = _visible_messages(messages)
    if len(visible) == 1:
        return visible[0].get("content", "")
    chunks = []
    for m in visible:
        chunks.append(
            f"## {m.get('role', 'assistant').upper()}\n{m.get('content', '')}"
        )
    return "\n\n".join(chunks)


def _extract_title(text, fallback):
    for line in text.splitlines():
        match = re.match(r"^\s*#\s+(.+?)\s*$", line)
        if match:
            return match.group(1).strip()[:100]
    return (fallback or "Local AI Document")[:100]


def _strip_first_h1(text):
    lines = text.splitlines()
    for i, line in enumerate(lines):
        if re.match(r"^\s*#\s+.+$", line):
            return "\n".join(lines[:i] + lines[i + 1:]).lstrip()
    return text


def _shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_text(cell, text, bold=False, color=TEXT, size=10.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _add_markdown(document, text, executive=False):
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            document.add_paragraph("")
            continue

        heading = re.match(r"^(#{1,6})\s+(.+)$", line)
        if heading:
            level = min(len(heading.group(1)), 3)
            p = document.add_paragraph()
            run = p.add_run(heading.group(2))
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(
                {1: 18, 2: 15, 3: 13}[level] + (1 if executive else 0)
            )
            run.font.color.rgb = RGBColor.from_string(
                DARK_RED if level == 1 else RED if level == 2 else TEXT
            )
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.space_before = Pt(8)
            continue

        bullet = re.match(r"^[-*]\s+(.+)$", line)
        numbered = re.match(r"^(\d+)\.\s+(.+)$", line)
        if bullet or numbered:
            text_value = bullet.group(1) if bullet else numbered.group(2)
            style = "List Bullet" if bullet else "List Number"
            p = document.add_paragraph(style=style)
            run = p.add_run(text_value)
            run.font.name = "Arial"
            run.font.size = Pt(12 if executive else 11.5)
            run.font.color.rgb = RGBColor.from_string(TEXT)
            continue

        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.25
        parts = re.split(r"(\*\*.+?\*\*)", line)
        for part in parts:
            if not part:
                continue
            bold = part.startswith("**") and part.endswith("**")
            clean = part[2:-2] if bold else part
            run = p.add_run(clean)
            run.bold = bold
            run.font.name = "Arial"
            run.font.size = Pt(12 if executive else 11.5)
            run.font.color.rgb = RGBColor.from_string(TEXT)


def _base_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11.5)
    return doc


def create_red_professional_docx(
    messages,
    title="Local AI Report",
    output_dir=OUTPUT_DIR,
):
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / (
        "local_ai_report_"
        + datetime.now().strftime("%Y%m%d_%H%M%S")
        + ".docx"
    )
    text = _document_text(messages)
    document_title = _extract_title(text, title)
    text = _strip_first_h1(text)

    doc = _base_document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    _shade_cell(cell, RED)
    _set_cell_text(cell, document_title, bold=True, color="FFFFFF", size=18)
    doc.add_paragraph("")
    _add_markdown(doc, text or "No document content.", executive=False)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = (
        "Red Professional · Generated "
        + datetime.now().strftime("%Y-%m-%d %H:%M")
    )
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(path)
    return path


def create_red_executive_docx(
    messages,
    title="Local AI Executive Report",
    output_dir=OUTPUT_DIR,
):
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / (
        "local_ai_executive_"
        + datetime.now().strftime("%Y%m%d_%H%M%S")
        + ".docx"
    )
    text = _document_text(messages)
    document_title = _extract_title(text, title)
    text = _strip_first_h1(text)

    doc = _base_document()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(document_title)
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string(TEXT)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("Local AI · Executive / Technical Report")
    r2.font.name = "Arial"
    r2.font.size = Pt(13)
    r2.font.color.rgb = RGBColor.from_string(MUTED)

    hero = doc.add_table(rows=2, cols=1)
    _shade_cell(hero.cell(0, 0), DARK_RED)
    _shade_cell(hero.cell(1, 0), DARK_RED)
    _set_cell_text(
        hero.cell(0, 0),
        "RED EXECUTIVE REPORT",
        bold=True,
        color="FFFFFF",
        size=17,
    )
    _set_cell_text(
        hero.cell(1, 0),
        "Structured local-model document output",
        color="FFE9EC",
        size=11.5,
    )

    doc.add_paragraph("")
    info = doc.add_table(rows=1, cols=3)
    for cell in info.rows[0].cells:
        _shade_cell(cell, LIGHT_GREY)
    _set_cell_text(
        info.cell(0, 0),
        "Preset\nRed Executive",
        bold=True,
        size=10.5,
    )
    _set_cell_text(
        info.cell(0, 1),
        "Execution\nLocal",
        bold=True,
        size=10.5,
    )
    _set_cell_text(
        info.cell(0, 2),
        "Generated\n" + datetime.now().strftime("%Y-%m-%d"),
        bold=True,
        size=10.5,
    )

    doc.add_paragraph("")
    summary = doc.add_table(rows=1, cols=1)
    _shade_cell(summary.cell(0, 0), LIGHT_RED)
    _set_cell_text(
        summary.cell(0, 0),
        "Executive Summary\n"
        "This document was generated locally using the selected AI model.",
        bold=True,
        color=TEXT,
        size=11.5,
    )

    doc.add_paragraph("")
    _add_markdown(doc, text or "No document content.", executive=True)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = (
        "Red Executive · Generated "
        + datetime.now().strftime("%Y-%m-%d %H:%M")
    )
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(path)
    return path
