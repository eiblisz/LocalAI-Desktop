from pathlib import Path

from openpyxl import load_workbook

from app.excel_tool import create_conversation_excel, create_structured_excel


def test_conversation_excel_created(tmp_path: Path):
    path = create_conversation_excel(
        [
            {"role": "user", "content": "hello"},
            {"role": "assistant", "content": "hi"},
        ],
        title="Conversation Test",
        output_dir=tmp_path,
    )
    wb = load_workbook(path)
    ws = wb["Conversation"]
    assert ws["A1"].value == "Conversation Test"
    assert ws["A3"].value == "Role"


def test_structured_excel_from_json(tmp_path: Path):
    payload = (
        '{"title":"Model Comparison","sheets":['
        '{"name":"Models","headers":["Model","Score"],'
        '"rows":[["Qwen",95],["Devstral",80]]}]}'
    )
    path = create_structured_excel(
        payload,
        output_dir=tmp_path,
    )
    wb = load_workbook(path)
    ws = wb["Models"]
    assert ws["A1"].value == "Model Comparison"
    assert ws["A4"].value == "Qwen"
    assert ws["B4"].value == 95
