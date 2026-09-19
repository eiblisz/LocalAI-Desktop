from pathlib import Path

from docx import Document
from openpyxl import Workbook

from app.internal_viewer import (
    classify_resource,
    render_docx_html,
    resource_title,
    spreadsheet_preview,
)


def test_resource_classifier_routes_web_and_supported_artifacts():
    assert classify_resource("https://example.com/report") == "browser"
    assert classify_resource(Path("report.html")) == "html"
    assert classify_resource(Path("report.pdf")) == "pdf"
    assert classify_resource(Path("report.docx")) == "docx"
    assert classify_resource(Path("report.xlsx")) == "spreadsheet"
    assert classify_resource(Path("report.csv")) == "spreadsheet"
    assert classify_resource(Path("report.md")) == "markdown"
    assert classify_resource(Path("report.txt")) == "text"
    assert classify_resource(Path("image.png")) == "image"


def test_resource_title_uses_host_or_file_name():
    assert resource_title("https://example.com/") == "example.com"
    assert resource_title("https://example.com/path/report") == "report"
    assert resource_title(Path("C:/tmp/report.pdf")) == "report.pdf"


def test_docx_preview_preserves_heading_paragraph_and_table(tmp_path):
    path = tmp_path / "report.docx"
    document = Document()
    document.add_heading("Executive Summary", level=1)
    document.add_paragraph("A concise document preview.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    document.save(path)

    rendered = render_docx_html(path)

    assert "<h1>Executive Summary</h1>" in rendered
    assert "A concise document preview." in rendered
    assert "Metric" in rendered
    assert "Value" in rendered


def test_spreadsheet_preview_reads_multiple_sheets(tmp_path):
    path = tmp_path / "book.xlsx"
    workbook = Workbook()
    first = workbook.active
    first.title = "Summary"
    first.append(["Name", "Value"])
    first.append(["BTC", 100])

    second = workbook.create_sheet("Notes")
    second.append(["Status"])
    second.append(["PASS"])
    workbook.save(path)

    preview = spreadsheet_preview(path)

    assert [item["name"] for item in preview] == ["Summary", "Notes"]
    assert preview[0]["rows"][1] == ["BTC", "100"]
    assert preview[1]["rows"][1] == ["PASS"]


def test_spreadsheet_preview_is_bounded(tmp_path):
    path = tmp_path / "large.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    for index in range(10):
        sheet.append([index, index + 1, index + 2])
    workbook.save(path)

    preview = spreadsheet_preview(path, max_rows=3, max_columns=2)

    assert len(preview[0]["rows"]) == 3
    assert len(preview[0]["rows"][0]) == 2
    assert preview[0]["truncated"] is True
