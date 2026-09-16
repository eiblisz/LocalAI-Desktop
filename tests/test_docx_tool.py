from pathlib import Path

from docx import Document

from app.docx_tool import (
    create_red_executive_docx,
    create_red_professional_docx,
)


def test_red_professional_docx_created(tmp_path: Path):
    path = create_red_professional_docx(
        [{"role": "assistant", "content": "# Test\n## Section\nBody text."}],
        title="Fallback",
        output_dir=tmp_path,
    )
    assert path.exists()
    assert path.suffix == ".docx"
    doc = Document(path)
    assert len(doc.paragraphs) >= 1


def test_red_executive_docx_created(tmp_path: Path):
    path = create_red_executive_docx(
        [{"role": "assistant", "content": "# Executive Test\nBody text."}],
        output_dir=tmp_path,
    )
    assert path.exists()
    assert path.name.startswith("local_ai_executive_")


def test_hungarian_red_executive_docx_localizes_fixed_labels(tmp_path: Path):
    path = create_red_executive_docx(
        [{
            "role": "assistant",
            "content": (
                "# Magyar helyi AI riport\n"
                "Ez egy magyar dokumentum a helyi modellekről és adatokról."
            ),
        }],
        output_dir=tmp_path,
    )
    doc = Document(path)
    table_text = "\n".join(
        cell.text
        for table in doc.tables
        for row in table.rows
        for cell in row.cells
    )
    assert "Vezetői összefoglaló" in table_text
    assert "Futtatás" in table_text
    assert "Készült" in table_text
